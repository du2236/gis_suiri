import arcpy
import shapefile
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import PatternFill
def detect_encoding(dbf_path):
    with open(dbf_path, "rb") as f:
        raw_data = f.read(10000)  
        result = chardet.detect(raw_data)
    return result['encoding']
def extract_shapefile_data(param_list):
    data = []
    m = 1
    for ly in param_list:
        shp_path = arcpy.Describe(ly).catalogPath
        dbf_path = shp_path.replace(".shp", ".dbf")
        cpg_path = shp_path.replace(".shp", ".cpg")
        
        if os.path.exists(cpg_path):
    	    with open(cpg_path, "r") as cpg_file:
    	        encoding = cpg_file.read().strip()
        else:
    	    encoding = detect_encoding(dbf_path)
            
        sr = shapefile.Reader(shp_path)
        for n, shape in enumerate(sr.shapes(), start=1):
            for i, (x, y) in enumerate(shape.points, start=1):
                data.append({"ID": int(f"{m}{n}"), "X": x, "Y": y})
        m += 1
    return pd.DataFrame(data)

def save_to_excel(df, output):
    wb = Workbook()
    ws = wb.active
    ws.title = "Shapefile Data"
    ws.append(["i","ID", "X", "Y"])  # ヘッダー
    
    for row in df.itertuples(index=True, name=None):
        ws.append(row)
    
    chart = ScatterChart()
    chart.legend = None
    chart.width = 16  # グラフのサイズ
    chart.height = 16
    chart.x_axis.tickLblPos = "low" 
    chart.y_axis.tickLblPos = "low"
    
    unique_ids = sorted(df["ID"].unique())
    color_palette = ["FF0000", "00FF00", "0000FF", "FFA500", "800080"]  # 赤, 緑, 青, オレンジ, 紫
    color_dict = {str(uid)[0]: color_palette[i % len(color_palette)] for i, uid in enumerate(unique_ids)}
    for i, uid in enumerate(unique_ids):
        rows = [r for r in range(2, len(df) + 2) if df.iloc[r-2]["ID"] == uid]
        if not rows:
            continue
        
        xvalues = Reference(ws, min_col=3, min_row=rows[0], max_row=rows[-1])
        yvalues = Reference(ws, min_col=4, min_row=rows[0], max_row=rows[-1])
        series = Series(yvalues, xvalues, title=str(uid))
        series.marker.symbol = "none"
        series.graphicalProperties.line.width = 45720  # ライン幅
        series.graphicalProperties.line.solidFill = color_dict[str(uid)[0]]
        chart.series.append(series)
    
    ws.add_chart(chart, "F2")
    wb.save(output)

def main():
    param_text = arcpy.GetParameterAsText(0)
    output = arcpy.GetParameterAsText(1)
    
    if param_text:
        param_list = param_text.split(";")
    else:
        arcpy.AddError("No input files provided.")
        return
    
    arcpy.SetProgressor("step", "Processing shapefiles...", 0, len(param_list), 1)
    df = extract_shapefile_data(param_list)
    save_to_excel(df, output)
    arcpy.ResetProgressor()

if __name__ == "__main__":
    main()
